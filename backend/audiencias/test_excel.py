import openpyxl
import pytest

from audiencias.excel import ABA, gerar_excel
from audiencias.relatorio import audiencias_do_mes
from audiencias.resumo import montar_resumo_excel


def planilha(ano, mes, sentencas=0):
    resumo = montar_resumo_excel(audiencias_do_mes(ano, mes))
    return openpyxl.load_workbook(gerar_excel(resumo, ano, mes, sentencas))


def test_so_a_aba_mes(dados):
    assert planilha(2026, 8).sheetnames == [ABA]


def test_cabecalho_e_rodape(dados):
    ws = planilha(2026, 8, sentencas=2)[ABA]
    assert ws["B2"].value.month == 8 and ws["B2"].value.year == 2026
    assert ws["D4"].value == 4      # agendadas
    assert ws["D5"].value == 3      # processuais
    assert ws["D6"].value == 1      # pré-processuais
    assert ws["D31"].value == 3     # realizadas (frutífera, infrutífera, frutífera)
    assert ws["D32"].value == 2     # acordos (só frutíferas)
    assert ws["D33"].value == 2     # sentenças digitadas


def test_bloco_pre(dados):
    ws = planilha(2026, 8)[ABA]
    assert ws["B10"].value == 1     # realizadas do bloco
    assert ws["D11"].value == 1     # realizadas na mediação
    assert ws["D12"].value == 1     # frutífera na mediação
    assert ws["E12"].value == 0


def test_bloco_processual(dados):
    ws = planilha(2026, 8)[ABA]
    assert ws["B21"].value == 2
    assert ws["D23"].value == 1     # 1ª Vara, mediação, frutífera
    assert ws["E24"].value == 1     # JEC, conciliação, infrutífera
    assert ws["E27"].value == 1     # 2ª Vara, conciliação, cancelada


def test_mes_sem_audiencias_gera_zeros(db):
    ws = planilha(2026, 1)[ABA]
    assert ws["D4"].value == 0
    assert ws["D23"].value == 0


def test_endpoint_baixa_o_excel(api, dados):
    r = api.get("/api/excel/?ano=2026&mes=8&sentencas=1")
    assert r.status_code == 200
    assert "spreadsheetml" in r["Content-Type"]
    ws = openpyxl.load_workbook(__import__("io").BytesIO(b"".join(r.streaming_content)))[ABA]
    assert ws["D4"].value == 4
    assert ws["D33"].value == 1


@pytest.mark.parametrize("consulta", ["", "?ano=2026", "?ano=2026&mes=13", "?ano=2026&mes=8&sentencas=-1", "?ano=x&mes=8"])
def test_endpoint_recusa_parametros_invalidos(api, consulta):
    assert api.get(f"/api/excel/{consulta}").status_code == 400