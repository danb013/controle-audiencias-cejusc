from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl

from .regras import Resultado

MODELO = Path(__file__).resolve().parent / "modelos" / "modelo_mensal.xlsx"
ABA = "MÊS"

# ---------------------------------------------------------------------------
# MAPA DAS CÉLULAS: se algum número cair no lugar errado, o ajuste é só aqui.
# ---------------------------------------------------------------------------
CELULA_MES = "B2"
CELULA_TOTAL_AGENDADAS = "D4"
CELULA_AGENDADAS_PROCESSUAL = "D5"
CELULA_AGENDADAS_PRE = "D6"

COLUNA_POR_TIPO = {"mediacao": "D", "conciliacao": "E"}

BLOCOS = {
    "pre": {
        "total_realizadas": "B10",   # frutíferas + infrutíferas do bloco
        "linha_realizadas": 11,      # realizadas de cada coluna (mediação / conciliação)
        "linhas": {
            Resultado.FRUTIFERA: 12,
            Resultado.INFRUTIFERA: 13,
            Resultado.REDESIGNADA: 14,
            Resultado.PREJ_AUS_REQUERIDO: 15,
            Resultado.CANCELADA: 16,
            Resultado.PREJ_AUS_PARTES: 17,
        },
    },
    "processual": {
        "total_realizadas": "B21",
        "linha_realizadas": 22,
        "linhas": {
            Resultado.FRUTIFERA: 23,
            Resultado.INFRUTIFERA: 24,
            Resultado.PREJ_AUS_REQUERIDO: 25,
            Resultado.PREJ_AUS_PARTES: 26,
            Resultado.CANCELADA: 27,
            Resultado.REDESIGNADA: 28,
        },
    },
}

CELULA_REALIZADAS = "D31"
CELULA_ACORDOS = "D32"
CELULA_SENTENCAS = "D33"


def gerar_excel(resumo: dict, ano: int, mes: int, sentencas_pre: int = 0) -> BytesIO:
    """Preenche a aba MÊS do modelo e devolve o arquivo em memória."""
    wb = openpyxl.load_workbook(MODELO)
    for nome in wb.sheetnames:
        if nome != ABA:
            del wb[nome]
    ws = wb[ABA]

    ws[CELULA_MES] = date(ano, mes, 1)
    ws[CELULA_TOTAL_AGENDADAS] = resumo["total_agendadas"]
    ws[CELULA_AGENDADAS_PROCESSUAL] = resumo["agendadas_processual"]
    ws[CELULA_AGENDADAS_PRE] = resumo["agendadas_pre"]

    for bloco, cfg in BLOCOS.items():
        dados = resumo["blocos"][bloco]
        realizadas_do_bloco = 0
        for tipo, coluna in COLUNA_POR_TIPO.items():
            for resultado, linha in cfg["linhas"].items():
                ws[f"{coluna}{linha}"] = dados[tipo][resultado.value]
            realizadas_do_tipo = (
                dados[tipo][Resultado.FRUTIFERA.value] + dados[tipo][Resultado.INFRUTIFERA.value]
            )
            ws[f"{coluna}{cfg['linha_realizadas']}"] = realizadas_do_tipo
            realizadas_do_bloco += realizadas_do_tipo
        ws[cfg["total_realizadas"]] = realizadas_do_bloco

    ws[CELULA_REALIZADAS] = resumo["realizadas"]
    ws[CELULA_ACORDOS] = resumo["acordos"]
    ws[CELULA_SENTENCAS] = sentencas_pre

    arquivo = BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)
    return arquivo